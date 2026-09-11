from __future__ import annotations

from calendar import monthrange
from collections import defaultdict
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from sqlalchemy import select

from database.db import async_session
from database.models import (
    DailyReport,
    FuelBatch,
    FuelConsumptionAllocation,
    User,
)


# ============================================================
# Константы
# ============================================================

MONTH_NAMES = (
    "январь",
    "февраль",
    "март",
    "апрель",
    "май",
    "июнь",
    "июль",
    "август",
    "сентябрь",
    "октябрь",
    "ноябрь",
    "декабрь",
)

THIN_SIDE = Side(style="thin", color="000000")

TABLE_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE,
)

TITLE_FONT = Font(
    name="Arial",
    size=14,
    bold=True,
)

HEADER_FONT = Font(
    name="Arial",
    size=10,
    bold=True,
)

DATA_FONT = Font(
    name="Arial",
    size=10,
)

CENTER = Alignment(
    horizontal="center",
    vertical="center",
    wrap_text=True,
)

LEFT = Alignment(
    horizontal="left",
    vertical="top",
    wrap_text=True,
)


# ============================================================
# Вспомогательные функции
# ============================================================

def _money(value: float | Decimal | None) -> float:
    """
    Округление денежных значений только при выводе в Excel.
    Внутренние расчёты не округляются.
    """
    if value is None:
        return 0.0

    result = Decimal(str(value)).quantize(
        Decimal("0.01"),
        rounding=ROUND_HALF_UP,
    )

    return float(result)


def _liters(value: float | Decimal | None) -> int:
    """
    Округление литровых итоговых значений до целого.
    """
    if value is None:
        return 0

    result = Decimal(str(value)).quantize(
        Decimal("1"),
        rounding=ROUND_HALF_UP,
    )

    return int(result)


def _period_dates(year: int, month: int) -> tuple[date, date]:
    if year < 1:
        raise ValueError("Некорректный год.")

    if not 1 <= month <= 12:
        raise ValueError("Месяц должен быть от 1 до 12.")

    first_day = date(year, month, 1)

    last_day = date(
        year,
        month,
        monthrange(year, month)[1],
    )

    return first_day, last_day


def _period_text(year: int, month: int) -> str:
    return f"{MONTH_NAMES[month - 1]} {year}"


def _safe_filename(value: str) -> str:
    forbidden = '<>:"/\\|?*'

    value = "".join(
        "_" if char in forbidden else char
        for char in value
    )

    value = " ".join(value.split()).strip()

    return value or "Оператор"


def _apply_border(
    worksheet,
    min_row: int,
    max_row: int,
    min_col: int,
    max_col: int,
) -> None:
    for row in worksheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            cell.border = TABLE_BORDER


# ============================================================
# Загрузка данных
# ============================================================

async def _load_month_data(
    user_id: int,
    first_day: date,
    last_day: date,
) -> tuple[
    User,
    list[DailyReport],
    list[DailyReport],
    list[FuelBatch],
    list[FuelConsumptionAllocation],
]:
    """
    Загружает данные для отчёта.

    month_reports:
        Только daily_reports выбранного месяца.

    historical_reports:
        Все daily_reports пользователя до конца выбранного месяца.
        Нужны для восстановления виртуального бака.

    batches:
        Все FIFO-партии пользователя до конца выбранного месяца.

    allocations:
        Все FIFO-распределения до конца выбранного месяца.

    БД только читается.
    """

    async with async_session() as session:

        # ----------------------------------------------------
        # Пользователь
        # ----------------------------------------------------

        user_result = await session.execute(
            select(User).where(
                User.id == user_id
            )
        )

        user = user_result.scalar_one_or_none()

        if user is None:
            raise ValueError(
                f"Пользователь с id={user_id} не найден."
            )

        # ----------------------------------------------------
        # Отчёты выбранного месяца
        # ----------------------------------------------------

        month_result = await session.execute(
            select(DailyReport)
            .where(
                DailyReport.user_id == user_id,
                DailyReport.report_date >= first_day,
                DailyReport.report_date <= last_day,
            )
            .order_by(
                DailyReport.report_date,
                DailyReport.id,
            )
        )

        month_reports = list(
            month_result.scalars().all()
        )

        # ----------------------------------------------------
        # Все исторические отчёты до конца месяца
        # ----------------------------------------------------

        historical_result = await session.execute(
            select(DailyReport)
            .where(
                DailyReport.user_id == user_id,
                DailyReport.report_date <= last_day,
            )
            .order_by(
                DailyReport.report_date,
                DailyReport.id,
            )
        )

        historical_reports = list(
            historical_result.scalars().all()
        )

        # ----------------------------------------------------
        # Все FIFO-партии до конца выбранного месяца
        # ----------------------------------------------------

        batches_result = await session.execute(
            select(FuelBatch)
            .where(
                FuelBatch.user_id == user_id,
                FuelBatch.daily_report_id.in_(
                    select(DailyReport.id)
                    .where(
                        DailyReport.user_id == user_id,
                        DailyReport.report_date <= last_day,
                    )
                ),
            )
            .order_by(FuelBatch.id)
        )

        batches = list(
            batches_result.scalars().all()
        )

        # ----------------------------------------------------
        # Все FIFO-распределения до конца месяца
        #
        # ВАЖНО:
        # allocation может быть создан значительно позже,
        # но относится к старому daily_report.
        #
        # Поэтому берём их по daily_report_id.
        # ----------------------------------------------------

        allocations_result = await session.execute(
            select(FuelConsumptionAllocation)
            .where(
                FuelConsumptionAllocation.daily_report_id.in_(
                    select(DailyReport.id)
                    .where(
                        DailyReport.user_id == user_id,
                        DailyReport.report_date <= last_day,
                    )
                )
            )
            .order_by(
                FuelConsumptionAllocation.id
            )
        )

        allocations = list(
            allocations_result.scalars().all()
        )

        return (
            user,
            month_reports,
            historical_reports,
            batches,
            allocations,
        )


# ============================================================
# Подготовка движения топлива
# ============================================================

def _build_fuel_movement(
    first_day: date,
    last_day: date,
    historical_reports: list[DailyReport],
    batches: list[FuelBatch],
    allocations: list[FuelConsumptionAllocation],
) -> list[dict]:
    """
    Формирует дневное движение топлива.

    FIFO НЕ рассчитывается.

    Используются уже сохранённые:
        fuel_batches
        fuel_consumption_allocations
        daily_reports.fuel_consumed

    История до начала выбранного месяца используется для
    определения остатка на начало месяца.

    День попадает в отчёт только если в нём есть движение
    топлива:
        - заправка;
        - расход топлива.
    """

    # --------------------------------------------------------
    # Индекс daily_report по ID
    # --------------------------------------------------------

    reports_by_id = {
        report.id: report
        for report in historical_reports
    }

    # --------------------------------------------------------
    # Движение полученного топлива.
    #
    # Источник:
    # FuelBatch.
    #
    # Это важно: fuel_liters/fuel_amount из daily_report
    # не используются для расчёта FIFO-остатка.
    # --------------------------------------------------------

    received_liters_by_date: dict[date, float] = defaultdict(float)
    received_cost_by_date: dict[date, float] = defaultdict(float)

    for batch in batches:
        report = reports_by_id.get(
            batch.daily_report_id
        )

        if report is None:
            continue

        batch_date = report.report_date

        received_liters_by_date[batch_date] += (
            batch.liters
        )

        received_cost_by_date[batch_date] += (
            batch.amount
        )

    # --------------------------------------------------------
    # Исторический расход.
    #
    # Источник:
    # DailyReport.fuel_consumed
    #
    # Никакого пересчёта через текущий расход пользователя.
    # --------------------------------------------------------

    consumed_liters_by_date: dict[date, float] = defaultdict(float)
    mileage_by_date: dict[date, float] = defaultdict(float)

    for report in historical_reports:
        report_date = report.report_date

        mileage_by_date[report_date] += (
            report.mileage or 0.0
        )

        consumed_liters_by_date[report_date] += (
            report.fuel_consumed or 0.0
        )

    # --------------------------------------------------------
    # Стоимость фактически распределённого расхода.
    #
    # Источник:
    # FuelConsumptionAllocation.cost
    #
    # Allocation относится к дате исходного daily_report.
    #
    # Это особенно важно при погашении отрицательного бака:
    # заправка сегодня может создать allocation для расхода
    # предыдущего дня.
    # --------------------------------------------------------

    consumed_cost_by_date: dict[date, float] = defaultdict(float)

    allocated_liters_by_date: dict[date, float] = defaultdict(float)

    for allocation in allocations:
        report = reports_by_id.get(
            allocation.daily_report_id
        )

        if report is None:
            continue

        report_date = report.report_date

        consumed_cost_by_date[report_date] += (
            allocation.cost
        )

        allocated_liters_by_date[report_date] += (
            allocation.liters
        )

    # --------------------------------------------------------
    # Определяем дни с движением топлива.
    #
    # Только выбранный месяц.
    # --------------------------------------------------------

    movement_dates = set()

    for report_date, liters in received_liters_by_date.items():
        if first_day <= report_date <= last_day:
            if liters != 0:
                movement_dates.add(report_date)

    for report_date, liters in consumed_liters_by_date.items():
        if first_day <= report_date <= last_day:
            if liters != 0:
                movement_dates.add(report_date)

    # --------------------------------------------------------
    # Если в выбранном месяце нет движения топлива,
    # лист будет содержать только заголовок.
    # --------------------------------------------------------

    if not movement_dates:
        return []

    # --------------------------------------------------------
    # Восстанавливаем состояние виртуального бака
    # последовательно с самого начала доступной истории.
    #
    # Это не FIFO-расчёт.
    #
    # FIFO уже сохранён в allocations.
    #
    # Здесь только суммирование:
    #
    # остаток литров =
    #     получено - исторически израсходовано
    #
    # остаток ДС =
    #     стоимость заправок - стоимость распределённого расхода
    # --------------------------------------------------------

    all_dates = set(received_liters_by_date)
    all_dates.update(consumed_liters_by_date)
    all_dates.update(consumed_cost_by_date)

    all_dates = sorted(
        report_date
        for report_date in all_dates
        if report_date <= last_day
    )

    liters_balance = 0.0
    cost_balance = 0.0

    result: list[dict] = []

    for current_date in all_dates:

        received_liters = received_liters_by_date.get(
            current_date,
            0.0,
        )

        received_cost = received_cost_by_date.get(
            current_date,
            0.0,
        )

        consumed_liters = consumed_liters_by_date.get(
            current_date,
            0.0,
        )

        consumed_cost = consumed_cost_by_date.get(
            current_date,
            0.0,
        )

        # Состояние ДО операций текущего дня.
        liters_start = liters_balance
        cost_start = cost_balance

        # Операции дня.
        liters_balance = (
            liters_balance
            + received_liters
            - consumed_liters
        )

        cost_balance = (
            cost_balance
            + received_cost
            - consumed_cost
        )

        liters_end = liters_balance
        cost_end = cost_balance

        # Нужна строка только для выбранного месяца
        # и только если это день движения топлива.
        if current_date not in movement_dates:
            continue

        mileage = mileage_by_date.get(
            current_date,
            0.0,
        )

        consumption_rate = None

        if mileage > 0:
            consumption_rate = (
                consumed_liters
                / mileage
                * 100
            )

        result.append(
            {
                "date": current_date,

                "mileage": mileage,

                "cost_start": cost_start,
                "received_cost": received_cost,
                "consumed_cost": consumed_cost,
                "cost_end": cost_end,

                "liters_start": liters_start,
                "received_liters": received_liters,
                "consumed_liters": consumed_liters,
                "liters_end": liters_end,

                "consumption_rate": consumption_rate,
            }
        )

    return result


# ============================================================
# Лист «Движение ДС и топлива»
# ============================================================

def _create_fuel_sheet(
    workbook: Workbook,
    user_name: str,
    period_text: str,
    movement: list[dict],
) -> None:

    worksheet = workbook.active

    worksheet.title = "Движение ДС и топлива"

    # --------------------------------------------------------
    # Заголовок
    # --------------------------------------------------------

    worksheet.merge_cells("A1:K1")

    title = worksheet["A1"]

    title.value = (
        f"Отчет по движению топлива "
        f"{user_name} за период: {period_text}"
    )

    title.font = TITLE_FONT
    title.alignment = CENTER

    # --------------------------------------------------------
    # Объединения заголовков
    # --------------------------------------------------------

    worksheet.merge_cells("A3:A4")
    worksheet.merge_cells("B3:B4")
    worksheet.merge_cells("C3:F3")
    worksheet.merge_cells("G3:J3")
    worksheet.merge_cells("K3:K4")

    worksheet["A3"] = "Дата"

    worksheet["B3"] = "Пройдено км"

    worksheet["C3"] = (
        "Остаток бензина в пересчете на ДС"
    )

    worksheet["G3"] = (
        "Остаток бензина в литрах"
    )

    worksheet["K3"] = (
        "Расчетный расход л/100км"
    )

    worksheet["C4"] = "На начало дня"
    worksheet["D4"] = "Получено"
    worksheet["E4"] = "Израсходовано"
    worksheet["F4"] = "На конец дня"

    worksheet["G4"] = "На начало дня"
    worksheet["H4"] = "Получено"
    worksheet["I4"] = "Израсходовано"
    worksheet["J4"] = "На конец дня"

    # --------------------------------------------------------
    # Формат заголовков
    # --------------------------------------------------------

    for row in worksheet.iter_rows(
        min_row=3,
        max_row=4,
        min_col=1,
        max_col=11,
    ):
        for cell in row:
            cell.font = HEADER_FONT
            cell.alignment = CENTER

    _apply_border(
        worksheet,
        3,
        4,
        1,
        11,
    )

    # --------------------------------------------------------
    # Данные
    # --------------------------------------------------------

    first_data_row = 5

    for row_number, item in enumerate(
        movement,
        start=first_data_row,
    ):

        worksheet.cell(
            row_number,
            1,
            item["date"],
        )

        worksheet.cell(
            row_number,
            2,
            item["mileage"],
        )

        worksheet.cell(
            row_number,
            3,
            _money(item["cost_start"]),
        )

        worksheet.cell(
            row_number,
            4,
            _money(item["received_cost"]),
        )

        worksheet.cell(
            row_number,
            5,
            _money(item["consumed_cost"]),
        )

        worksheet.cell(
            row_number,
            6,
            _money(item["cost_end"]),
        )

        worksheet.cell(
            row_number,
            7,
            _liters(item["liters_start"]),
        )

        worksheet.cell(
            row_number,
            8,
            _liters(item["received_liters"]),
        )

        worksheet.cell(
            row_number,
            9,
            _liters(item["consumed_liters"]),
        )

        worksheet.cell(
            row_number,
            10,
            _liters(item["liters_end"]),
        )

        worksheet.cell(
            row_number,
            11,
            (
                item["consumption_rate"]
                if item["consumption_rate"] is not None
                else None
            ),
        )

    # --------------------------------------------------------
    # Формат данных
    # --------------------------------------------------------

    if movement:

        last_data_row = (
            first_data_row
            + len(movement)
            - 1
        )

        for row in worksheet.iter_rows(
            min_row=first_data_row,
            max_row=last_data_row,
            min_col=1,
            max_col=11,
        ):
            for cell in row:
                cell.font = DATA_FONT
                cell.alignment = CENTER

        _apply_border(
            worksheet,
            first_data_row,
            last_data_row,
            1,
            11,
        )

        for row_number in range(
            first_data_row,
            last_data_row + 1,
        ):

            worksheet.cell(
                row_number,
                1,
            ).number_format = "DD.MM.YYYY"

            worksheet.cell(
                row_number,
                2,
            ).number_format = "#,##0.00"

            for column in range(3, 7):
                worksheet.cell(
                    row_number,
                    column,
                ).number_format = "#,##0.00"

            for column in range(7, 11):
                worksheet.cell(
                    row_number,
                    column,
                ).number_format = "0"

            worksheet.cell(
                row_number,
                11,
            ).number_format = "0.00"

    # --------------------------------------------------------
    # Размеры
    # --------------------------------------------------------

    widths = {
        "A": 13,
        "B": 14,
        "C": 17,
        "D": 15,
        "E": 17,
        "F": 17,
        "G": 15,
        "H": 14,
        "I": 16,
        "J": 15,
        "K": 18,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width

    worksheet.row_dimensions[1].height = 25
    worksheet.row_dimensions[3].height = 42
    worksheet.row_dimensions[4].height = 42

    worksheet.freeze_panes = "A5"


# ============================================================
# Лист «Отчет о выполненных работах»
# ============================================================

def _create_work_sheet(
    workbook: Workbook,
    user_name: str,
    period_text: str,
    reports: list[DailyReport],
) -> None:

    worksheet = workbook.create_sheet(
        "Отчет о выполненных работах"
    )

    # --------------------------------------------------------
    # Заголовок
    # --------------------------------------------------------

    worksheet.merge_cells("A1:C1")

    title = worksheet["A1"]

    title.value = (
        f"Отчет по выполненным работам "
        f"{user_name} за период: {period_text}"
    )

    title.font = TITLE_FONT
    title.alignment = CENTER

    # --------------------------------------------------------
    # Заголовки
    # --------------------------------------------------------

    worksheet["A3"] = "Дата"
    worksheet["B3"] = "Место проведения работ"
    worksheet["C3"] = "Выполненные работы"

    for cell in worksheet[3]:
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = TABLE_BORDER

    # --------------------------------------------------------
    # Данные
    # --------------------------------------------------------

    first_data_row = 4

    for row_number, report in enumerate(
        reports,
        start=first_data_row,
    ):

        worksheet.cell(
            row_number,
            1,
            report.report_date,
        )

        worksheet.cell(
            row_number,
            2,
            report.workplace,
        )

        worksheet.cell(
            row_number,
            3,
            report.work_description,
        )

    # --------------------------------------------------------
    # Формат данных
    # --------------------------------------------------------

    if reports:

        last_data_row = (
            first_data_row
            + len(reports)
            - 1
        )

        for row in worksheet.iter_rows(
            min_row=first_data_row,
            max_row=last_data_row,
            min_col=1,
            max_col=3,
        ):
            for cell in row:
                cell.font = DATA_FONT
                cell.border = TABLE_BORDER

        for row_number in range(
            first_data_row,
            last_data_row + 1,
        ):

            worksheet.cell(
                row_number,
                1,
            ).alignment = CENTER

            worksheet.cell(
                row_number,
                1,
            ).number_format = "DD.MM.YYYY"

            worksheet.cell(
                row_number,
                2,
            ).alignment = LEFT

            worksheet.cell(
                row_number,
                3,
            ).alignment = LEFT

    # --------------------------------------------------------
    # Размеры
    # --------------------------------------------------------

    worksheet.column_dimensions["A"].width = 13
    worksheet.column_dimensions["B"].width = 30
    worksheet.column_dimensions["C"].width = 70

    worksheet.row_dimensions[1].height = 25
    worksheet.row_dimensions[3].height = 35

    worksheet.freeze_panes = "A4"


# ============================================================
# Публичная функция
# ============================================================

async def generate_monthly_report(
    user_id: int,
    year: int,
    month: int,
) -> tuple[BytesIO, str]:
    """
    Формирует месячный Excel-отчёт конкретного оператора.

    Возвращает:
        BytesIO с готовым XLSX
        имя файла

    БД:
        только SELECT.
        INSERT / UPDATE / DELETE отсутствуют.
    """

    first_day, last_day = _period_dates(
        year,
        month,
    )

    (
        user,
        month_reports,
        historical_reports,
        batches,
        allocations,
    ) = await _load_month_data(
        user_id=user_id,
        first_day=first_day,
        last_day=last_day,
    )

    # --------------------------------------------------------
    # Формируем движение топлива.
    # --------------------------------------------------------

    movement = _build_fuel_movement(
        first_day=first_day,
        last_day=last_day,
        historical_reports=historical_reports,
        batches=batches,
        allocations=allocations,
    )

    # --------------------------------------------------------
    # Создаём книгу.
    # --------------------------------------------------------

    workbook = Workbook()

    period_text = _period_text(
        year,
        month,
    )

    # Первый лист.
    _create_fuel_sheet(
        workbook=workbook,
        user_name=user.name,
        period_text=period_text,
        movement=movement,
    )

    # Второй лист.
    _create_work_sheet(
        workbook=workbook,
        user_name=user.name,
        period_text=period_text,
        reports=month_reports,
    )

    # --------------------------------------------------------
    # Сохраняем в память.
    # --------------------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    # --------------------------------------------------------
    # Имя файла.
    # --------------------------------------------------------

    safe_name = _safe_filename(
        user.name
    )

    filename = (
        f"Отчет_{safe_name}_"
        f"{year:04d}-{month:02d}.xlsx"
    )

    return output, filename


# ============================================================
# Дополнительная функция локального сохранения
# ============================================================

async def save_monthly_report(
    user_id: int,
    year: int,
    month: int,
    directory: str | Path,
) -> Path:
    """
    Сохраняет готовый отчёт на диск.

    В Telegram эта функция не обязательна:
    бот может отправить BytesIO напрямую.
    """

    output, filename = await generate_monthly_report(
        user_id=user_id,
        year=year,
        month=month,
    )

    directory = Path(directory)

    directory.mkdir(
        parents=True,
        exist_ok=True,
    )

    path = directory / filename

    path.write_bytes(
        output.getvalue()
    )

    return path